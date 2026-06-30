"""Regression: export filenames use the correct extension.

Bug (2026-06-30): the Excel export wrote files named ``*.excel`` instead of
``*.xlsx`` because the format key ("excel") was used directly as the file
extension. JSON/XML happened to be correct only because their format key
equals their extension.
"""

from __future__ import annotations

from plck.export import export_tags
from plck.models import PlcConnection, TagDataType, TagNode, TagTree


def _tree() -> TagTree:
    return TagTree(
        plc=PlcConnection(nickname="Compressor Skid A", ip="192.0.2.10"),
        root=[TagNode(name="Speed", data_type=TagDataType.DINT)],
        discovered_at="2026-06-30T10:15:00Z",
    )


def _udt_tree() -> TagTree:
    udt = TagNode(name="PIT204", data_type=TagDataType.UDT, udt_name="PRESSURE")
    udt.members = [
        TagNode(name="PIT204.PV", data_type=TagDataType.REAL),
        TagNode(name="PIT204.SP", data_type=TagDataType.REAL),
    ]
    return TagTree(
        plc=PlcConnection(nickname="Skid", ip="192.0.2.10"),
        root=[TagNode(name="Speed", data_type=TagDataType.DINT), udt],
        discovered_at="2026-06-30T10:15:00Z",
    )


def test_excel_export_uses_xlsx_extension(tmp_path):
    written = export_tags(_tree(), fmt="excel", out_dir=tmp_path)
    assert len(written) == 1
    assert written[0].suffix == ".xlsx", f"expected .xlsx, got {written[0].name}"
    assert written[0].is_file()


def test_excel_is_single_flattened_tags_sheet(tmp_path):
    from openpyxl import load_workbook

    written = export_tags(_udt_tree(), fmt="excel", out_dir=tmp_path)[0]
    wb = load_workbook(written)
    # One workbook: a PLC metadata sheet + a single 'Tags' sheet — NOT one
    # sheet per UDT.
    assert wb.sheetnames == ["PLC", "Tags"]
    ws = wb["Tags"]
    headers = [c.value for c in ws[1]]
    assert headers[:3] == ["Level", "Tag", "Data Type"]
    # All tags flattened into this one sheet, in tree order with depth.
    rows = [(r[0].value, r[1].value) for r in ws.iter_rows(min_row=2)]
    tags = [name for _lvl, name in rows]
    assert tags == ["Speed", "PIT204", "PIT204.PV", "PIT204.SP"]
    levels = {name: lvl for lvl, name in rows}
    assert levels["PIT204"] == 0 and levels["PIT204.PV"] == 1


def test_json_export_uses_json_extension(tmp_path):
    written = export_tags(_tree(), fmt="json", out_dir=tmp_path)
    assert written[0].suffix == ".json"
    assert written[0].is_file()


def test_xml_export_uses_xml_extension(tmp_path):
    written = export_tags(_tree(), fmt="xml", out_dir=tmp_path)
    assert written[0].suffix == ".xml"
    assert written[0].is_file()
