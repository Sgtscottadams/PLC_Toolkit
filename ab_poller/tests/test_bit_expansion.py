"""Bit expansion: integer tags expose individual BOOL bits (Tag.0 .. Tag.N).

Mirrors Studio 5000 bit addressing for alarm/status words. Per-bit
descriptions are NOT read (pylogix can't surface tag comments over CIP), so
bits are unnamed BOOLs. Bit values are derived from the parent word, not read
per-bit.
"""

from __future__ import annotations

from plck.discovery import _build_node
from plck.models import PlcConnection, TagDataType, TagNode, TagTree


def test_dint_scalar_expands_to_32_bits():
    node = _build_node("Well_1.WELL_ALARM_WORD", "DINT", {}, max_depth=8, depth=0)
    assert node.data_type == TagDataType.DINT
    assert node.is_bit_word()
    assert len(node.members) == 32
    assert node.members[0].name == "Well_1.WELL_ALARM_WORD.0"
    assert node.members[31].name == "Well_1.WELL_ALARM_WORD.31"
    assert all(m.data_type == TagDataType.BOOL for m in node.members)
    assert all(m.bit_index == i for i, m in enumerate(node.members))


def test_bool_and_real_are_not_expanded():
    assert _build_node("Flag", "BOOL", {}, max_depth=8, depth=0).members == []
    assert _build_node("PV", "REAL", {}, max_depth=8, depth=0).members == []


def test_leaf_names_reads_word_not_bits():
    word = _build_node("Tank.ALARMS", "INT", {}, max_depth=8, depth=0)
    tree = TagTree(
        plc=PlcConnection(nickname="T", ip="192.0.2.5"),
        root=[word],
        discovered_at="2026-06-30T00:00:00Z",
    )
    # The post-discovery batch read targets the WORD (one read), never the
    # 16 derived bits.
    assert tree.leaf_names() == ["Tank.ALARMS"]
    # The word counts as a single leaf.
    assert tree.leaf_count() == 1
    # But the bits still exist in the tree for browse/export.
    assert tree.count() == 1 + 16


def test_exports_exclude_bits(tmp_path):
    from openpyxl import load_workbook

    from plck.export import export_tags

    word = _build_node("Tank.ALARMS", "DINT", {}, max_depth=8, depth=0)
    tree = TagTree(
        plc=PlcConnection(nickname="T", ip="192.0.2.5"),
        root=[word],
        discovered_at="2026-06-30T00:00:00Z",
    )

    # JSON: word present, no bit members.
    import json as _json
    jpath = export_tags(tree, fmt="json", out_dir=tmp_path)[0]
    payload = _json.loads(jpath.read_text(encoding="utf-8"))
    assert payload["tags"][0]["name"] == "Tank.ALARMS"
    assert "members" not in payload["tags"][0]

    # XML: only the word element, no bit children.
    xpath = export_tags(tree, fmt="xml", out_dir=tmp_path)[0]
    assert xpath.read_text(encoding="utf-8").count("<Tag ") == 1

    # Excel: one tag row (plus header), bits excluded.
    epath = export_tags(tree, fmt="excel", out_dir=tmp_path)[0]
    ws = load_workbook(epath)["Tags"]
    tag_rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[1]]
    assert [r[1] for r in tag_rows] == ["Tank.ALARMS"]


def test_db_snapshot_excludes_bits_but_reexpands_on_load(tmp_path):
    from plck.db import DatabaseManager, _serialize_tree

    word = _build_node("Tank.ALARMS", "DINT", {}, max_depth=8, depth=0)
    tree = TagTree(
        plc=PlcConnection(nickname="T", ip="192.0.2.5"),
        root=[word],
        discovered_at="2026-06-30T00:00:00Z",
    )

    # Serialized form carries the word but NOT its 32 bits.
    payload = _serialize_tree(tree)
    assert payload[0]["name"] == "Tank.ALARMS"
    assert "members" not in payload[0]

    # Round-trip through the DB re-derives the bits so a loaded scan matches
    # a live discovery.
    db = DatabaseManager(path=str(tmp_path / "t.db"))
    plc_id = db.upsert_plc(tree.plc)
    scan_id = db.save_tag_snapshot(plc_id, tree)
    loaded = db.load_scan(scan_id)
    assert loaded is not None
    assert loaded.root[0].is_bit_word()
    assert len(loaded.root[0].members) == 32
    assert loaded.root[0].members[5].bit_index == 5


def test_bit_value_derivation():
    # bit i of the word value = (word >> i) & 1
    word_value = 0b1010  # bits 1 and 3 set
    bits = [(word_value >> i) & 1 for i in range(4)]
    assert bits == [0, 1, 0, 1]
